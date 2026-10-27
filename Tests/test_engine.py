"""Integration suite: real decoders inspect outputs, all fixtures stay in a temporary folder."""
import sys,os,tempfile,json,subprocess,hashlib,io,zipfile,wave,math,struct
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
RES=ROOT/'Orbita.app/Contents/Resources'
sys.path.insert(0,str(RES/'Engine'))
import engine as E
from PIL import Image,ImageDraw,ImageFont
import pymupdf as fitz
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from catalog import A

passed=[];failed=[]
suite=os.environ.get("ORBITA_TEST_SUITE","all")
with tempfile.TemporaryDirectory(prefix='orbita-integration-') as temp:
 T=Path(temp)
 im=Image.new('RGBA',(480,320),(20,90,180,255));d=ImageDraw.Draw(im);d.text((30,80),'ORBITA TEST 123',font=E.font(28),fill='white');d.rectangle((100,150,250,250),fill='orange');png=T/'sample.png';im.save(png)
 frames=[im,Image.new('RGBA',im.size,'green')];gif=T/'animated.gif';frames[0].save(gif,save_all=True,append_images=frames[1:],duration=[100,200],loop=0)
 pdf=T/'sample.pdf';doc=fitz.open()
 for i in range(3):
  pg=doc.new_page(width=400,height=300);pg.insert_text((40,40),'SECRET ORBITA PAGE '+str(i+1));pg.insert_image(fitz.Rect(40,60,250,200),filename=png)
 doc.save(pdf);doc.close()
 tablepdf=T/'table.pdf';doc=fitz.open();pg=doc.new_page()
 for y in [50,80,110,140]:pg.draw_line((50,y),(350,y))
 for x in [50,200,350]:pg.draw_line((x,50),(x,140))
 for i,row in enumerate([['Name','Value'],['Orange','12'],['Blue','24']]):
  for j,value in enumerate(row):pg.insert_text((60+j*150,70+i*30),value)
 doc.save(tablepdf);doc.close()
 formpdf=T/'form.pdf';doc=fitz.open();pg=doc.new_page();w=fitz.Widget();w.field_name='Nome';w.field_type=fitz.PDF_WIDGET_TYPE_TEXT;w.rect=fitz.Rect(30,30,250,70);pg.add_widget(w);doc.save(formpdf);doc.close()
 text=T/'sample.md';text.write_text('# Orbita\n\nTesto **formattato** e una lista.\n\n- Arancia\n- Blu\n')
 word=T/'sample.docx';doc=Document();doc.add_heading('Orbita',0);doc.add_paragraph('Documento di prova.');doc.save(word)
 slide=T/'slides.pptx';prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[0]);s.shapes.title.text='Orbita';prs.save(slide)
 csv=T/'data.csv';csv.write_text('Name,Value\nOrange,12\nBlue,24\n')
 xlsx=T/'data.xlsx';book=Workbook();book.active.append(['Name','Value']);book.active.append(['Orange',12]);book.save(xlsx)
 srt=T/'subtitle.srt';srt.write_text('1\n00:00:00,000 --> 00:00:01,000\nOrbita test\n')
 audio=T/'sound.wav'
 with wave.open(str(audio),'wb') as w:
  w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b''.join(struct.pack('<h',int(10000*math.sin(2*math.pi*440*i/16000))) for i in range(32000)))
 video=T/'video.mp4';E.ff(['-f','lavfi','-i','color=c=blue:s=320x240:r=15:d=2','-i',audio,'-shortest']+E.vcodec('mp4')+[video])
 archive=T/'archive.zip'
 with zipfile.ZipFile(archive,'w') as z:z.writestr('hello.txt','Orbita')
 hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in [png,pdf,audio,video,word,csv,xlsx,srt]}
 def test(id,paths=None,params=None,check=None):
  office = id.startswith('slides.') or id=='sheet.to.pdf' or (id.startswith('document.') and paths and Path(paths[0]).suffix=='.docx')
  if suite=='core' and office:return []
  if suite=='office' and not office:return []
  E.OUTPUTS.clear()
  try:
   action=next(a for a in A if a['id']==id)
   defaults={f['key']:f['default'] for f in action['fields']};defaults.update(params or {})
   out=E.execute({'action':id,'paths':[str(x) for x in (paths or [png])],'params':defaults})
   assert out,'No output';assert all(Path(x).exists() for x in out)
   for item in out:
    p=Path(item)
    if p.is_file():assert p.stat().st_size>0
    if p.suffix in ['.png','.jpg','.gif','.webp','.avif','.heic','.tiff','.bmp','.ico','.icns']:Image.open(p).load()
    if p.suffix=='.pdf':
     with fitz.open(p) as doc:
      if not doc.needs_pass:assert len(doc)>0
   if check:check(out)
   passed.append(id);print('PASS',id,flush=True)
   return list(out)
  except Exception as e:failed.append({'action':id,'error':str(e)});print('FAIL',id,str(e)[:800],flush=True);return []
 for a in A:
  id=a['id']
  if id.startswith('image.'):
   if id in ['image.combine','image.collage','image.animate']:test(id,[png,png])
   else:test(id,params={'width':'120','height':'100','x':'10','y':'10','mb':'.02','text':'Orbita'})
  elif id=='vision.qr':
   qr=T/'qr.png';subprocess.run(['xcrun','swift',str(ROOT/'Tests/qr_fixture.swift'),str(qr)],check=True);test(id,[qr],check=lambda out: (_ for _ in ()).throw(AssertionError('QR payload mismatch')) if 'Orbita QR verification 2026' not in Path(out[0]).read_text() else None)
  elif id.startswith('vision.'):test(id)
 for a in A:
  id=a['id']
  if id.startswith('pdf.'):
   paths=[tablepdf] if id=='pdf.to.xlsx' else [formpdf] if id in ['pdf.forms','pdf.fill'] else [pdf,pdf] if id=='pdf.merge' else [pdf]
   p={'text':'SECRET','password':'test-password','file':str(png),'values':'{"Nome":"Mario"}','width':'100','height':'50','x':'30','y':'30'}
   if id=='pdf.unlock':
    encrypted=test('pdf.password',[pdf],{'password':'test-password'})
    if encrypted:paths=[Path(encrypted[0])]
   test(id,paths,p)
 for a in A:
  id=a['id']
  if id.startswith(('audio.','video.')) and id not in ['video.extractsubs','video.autosubs']:
   path=video if id.startswith('video.') else audio
   paths=[path,path] if id.endswith('.join') else [path]
   test(id,paths,{'start':'.2','end':'1.2','fade':'.2','width':'100','height':'100','x':'10','y':'10','file':str(srt if id.startswith('video.') else png),'mb':'.1'})
 for a in A:
  id=a['id']
  if id.startswith('document.to.'):test(id,[text])
  elif id.startswith('slides.to.'):test(id,[slide])
  elif id.startswith('sheet.to.'):test(id,[csv])
  elif id.startswith('subtitle.'):test(id,[srt])
  elif id.startswith('archive.'):test(id,[archive] if id.endswith('.extract') else [png,text],{'password':'test-password'})
  elif id=='files.renameOriginals':
   rename_sample=T/'rename-copy.png';rename_sample.write_bytes(png.read_bytes());test(id,[rename_sample])
  elif id.startswith('files.'):test(id,[png])
 test('document.to.pdf',[word])
 test('document.to.odt',[word])
 # Regression checks: frames, actual redaction, original integrity, collision and path traversal.
 def animation(out):assert Image.open(out[0]).n_frames==2
 test('image.to.webp',[gif],check=animation)
 def redact(out):
  with fitz.open(out[0]) as doc:assert 'SECRET' not in ''.join(p.get_text() for p in doc)
 test('pdf.redact',[pdf],{'text':'SECRET'},redact)
 first=test('image.to.jpg');second=test('image.to.jpg');assert suite=="office" or first!=second
 malicious=T/'malicious.zip'
 with zipfile.ZipFile(malicious,'w') as z:z.writestr('../escape.txt','no')
 try:E.execute({'action':'archive.extract','paths':[str(malicious)],'params':{}});raise AssertionError('Traversal accepted')
 except E.UserError:passed.append('security.archive-traversal');print('PASS security.archive-traversal')
 assert not (T/'escape.txt').exists()
 for path,digest in hashes.items():assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==digest
 passed.append('originals-unchanged');print('PASS originals-unchanged')
 report={'passed':passed,'failed':failed,'count':len(passed),'fixtures':'generated locally; deleted after test'}
 reportPath=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'TEST-RESULTS.json';reportPath.write_text(json.dumps(report,ensure_ascii=False,indent=2))
 print(json.dumps({'passed':len(passed),'failed':len(failed)}));sys.exit(bool(failed))
