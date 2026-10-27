#!/usr/bin/env python3
"""Orbita local engine. Requests and responses are JSON; originals are never overwritten."""
import sys
sys.dont_write_bytecode=True
import os,platform,json,math,re,io,csv,zipfile,tarfile,gzip,bz2,lzma,hashlib,shutil,tempfile,subprocess,html,time
from pathlib import Path
from contextlib import contextmanager
ROOT=Path(__file__).resolve().parent
RES=ROOT.parent
os.environ['PYTHONDONTWRITEBYTECODE']='1'
from PIL import Image,ImageOps,ImageDraw,ImageFont,ImageEnhance,ImageFilter,ImageSequence,ImageChops,ImageCms
import pillow_heif
pillow_heif.register_heif_opener()
Image.MAX_IMAGE_PIXELS=120_000_000
OUTPUTS=[]
DESTINATIONS={}
class UserError(Exception):pass

def run(args,timeout=3600,input=None):
 p=subprocess.run([str(x) for x in args],input=input,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout)
 if p.returncode: raise UserError(p.stderr.decode(errors='replace')[-3500:] or p.stdout.decode(errors='replace')[-2000:] or 'Operazione fallita.')
 return p.stdout

def number(p,k,default,lo=None,hi=None):
 try:v=float(p.get(k,default))
 except (ValueError,TypeError):raise UserError('Valore numerico non valido: '+k)
 if not math.isfinite(v) or (lo is not None and v<lo) or (hi is not None and v>hi):raise UserError('Valore fuori intervallo: '+k)
 return v

def nonempty(p,k):
 s=str(p.get(k,''))
 if not s:raise UserError('Compila il campo '+k)
 return s

def output(data,source,ext,suffix='-orbita'):
 source=Path(DESTINATIONS.get(str(source),source));stem=source.stem+suffix
 for i in range(10000):
  dest=source.parent/(stem+(f'-{i}' if i else '')+('.'+ext if ext else ''))
  try:
   with dest.open('xb') as f:
    try:f.write(data)
    except BaseException:
     f.close();dest.unlink(missing_ok=True);raise
   OUTPUTS.append(str(dest));return dest
  except FileExistsError:continue
 raise UserError('Troppi risultati con lo stesso nome.')

def publish(path,source,ext,suffix='-orbita'):
 # Stream large results while atomically reserving an unused name.
 source=Path(DESTINATIONS.get(str(source),source))
 for i in range(10000):
  dest=source.parent/(source.stem+suffix+(f'-{i}' if i else '')+('.'+ext if ext else ''))
  try:
   with dest.open('xb') as f:
    try:
     with open(path,'rb') as src:shutil.copyfileobj(src,f,1024*1024)
    except BaseException:
     f.close();dest.unlink(missing_ok=True);raise
   OUTPUTS.append(str(dest));return dest
  except FileExistsError:continue
 raise UserError('Troppi risultati con lo stesso nome.')

def folder(source,suffix):
 source=Path(DESTINATIONS.get(str(source),source))
 for i in range(10000):
  dest=source.parent/(source.stem+suffix+(f'-{i}' if i else ''))
  try:dest.mkdir();OUTPUTS.append(str(dest));return dest
  except FileExistsError:continue
 raise UserError('Impossibile creare cartella risultati.')

def image(source):
 with Image.open(source) as src:
  src.load();return color_managed(ImageOps.exif_transpose(src).copy())

def color_managed(im):
 profile=im.info.get('icc_profile')
 if not profile:return im
 try:
  alpha=im.convert('RGBA').getchannel('A') if 'A' in im.getbands() or 'transparency' in im.info else None
  source=im if im.mode in ['RGB','CMYK','L','LAB'] else im.convert('RGB')
  target=ImageCms.ImageCmsProfile(ImageCms.createProfile('sRGB'))
  converted=ImageCms.profileToProfile(source,ImageCms.ImageCmsProfile(io.BytesIO(profile)),target,outputMode='RGB')
  converted.info.update(im.info);converted.info['icc_profile']=target.tobytes()
  if alpha is not None:converted.putalpha(alpha)
  return converted
 except Exception as e:raise UserError('Profilo colore non leggibile: '+str(e))

def animation_loop(src,ext):
 loop=src.info.get('loop')
 if ext=='gif':
  if src.format=='WEBP':return {} if loop==1 else {'loop':0 if loop==0 else loop-1}
  return {} if loop is None else {'loop':loop}
 if ext=='webp':return {'loop':(1 if loop is None else 0 if loop==0 else loop+1) if src.format=='GIF' else (loop or 0)}
 return {}

def font(size):
 for path in ['/System/Library/Fonts/Supplemental/Arial.ttf','/System/Library/Fonts/Helvetica.ttc']:
  if Path(path).exists():return ImageFont.truetype(path,int(size))
 return ImageFont.load_default(size=int(size))

def flatten_white(im):
 rgba=im.convert('RGBA');bg=Image.new('RGBA',rgba.size,'white');bg.alpha_composite(rgba);return bg.convert('RGB')

def image_bytes(im,ext,quality=92,**kwargs):
 profile=im.info.get('icc_profile')
 if profile and ext in ['jpg','png','webp','tiff','heic','avif']:kwargs.setdefault('icc_profile',profile)
 fmt={'jpg':'JPEG','heic':'HEIF','tif':'TIFF'}.get(ext,ext.upper())
 if ext in ['jpg','bmp','pdf']:im=flatten_white(im)
 if ext=='ico':im=ImageOps.contain(im,(256,256));kwargs['sizes']=[(n,n) for n in [16,32,48,64,128,256] if n<=max(im.size)]
 if ext=='icns':
  canvas=Image.new('RGBA',(1024,1024));im=ImageOps.contain(im.convert('RGBA'),canvas.size);canvas.alpha_composite(im,((1024-im.width)//2,(1024-im.height)//2));im=canvas
 b=io.BytesIO();im.save(b,format=fmt,quality=quality,**kwargs);return b.getvalue()

def mask_subject(im):
 import numpy as np,onnxruntime as ort
 model=RES/'Models/u2netp.onnx'
 if not model.exists():raise UserError('Modello rimozione sfondo mancante.')
 opts=ort.SessionOptions();opts.intra_op_num_threads=max(1,min(4,os.cpu_count() or 2))
 sess=ort.InferenceSession(str(model),sess_options=opts,providers=['CPUExecutionProvider'])
 arr=np.asarray(im.convert('RGB').resize((320,320)),dtype=np.float32)/255
 arr=(arr-np.array([.485,.456,.406],dtype=np.float32))/np.array([.229,.224,.225],dtype=np.float32)
 pred=sess.run(None,{sess.get_inputs()[0].name:arr.transpose(2,0,1)[None]})[0][0,0]
 pred=(pred-pred.min())/(pred.max()-pred.min()+1e-8)
 return Image.fromarray((pred*255).astype('uint8')).resize(im.size,Image.Resampling.LANCZOS)

def vision(source,mode='ocr'):
 helper=RES/'bin/orbita-vision'
 if not helper.exists():raise UserError('Motore Vision mancante.')
 return json.loads(run([helper,mode,source],timeout=300))

def process_image(action,paths,p):
 op=action.split('.',1)[1]
 if op in ['combine','collage','animate']:
  ims=[image(x).convert('RGBA') for x in paths]
  if op=='combine':
   output(image_bytes(ims[0],'pdf',save_all=True,append_images=[flatten_white(x) for x in ims[1:]]),paths[0],'pdf');return
  if op=='collage':
   cols=int(number(p,'columns',2,1,20));cell=600;canvas=Image.new('RGB',(cell*cols,cell*math.ceil(len(ims)/cols)),'white')
   for i,im in enumerate(ims):im=ImageOps.contain(im,(cell-20,cell-20));canvas.paste(im,((i%cols)*cell+(cell-im.width)//2,(i//cols)*cell+(cell-im.height)//2),im)
   output(image_bytes(canvas,'png'),paths[0],'png');return
  size=ims[0].size;frames=[ImageOps.pad(im,size,color=(255,255,255,0)) for im in ims]
  output(image_bytes(frames[0],'gif',save_all=True,append_images=frames[1:],duration=int(number(p,'duration',250,20,60000)),loop=0,disposal=2),paths[0],'gif');return
 for path in paths:
  if op=='metadata':
   with Image.open(path) as raw:output(json.dumps({'size':raw.size,'mode':raw.mode,'info':raw.info,'exif':dict(raw.getexif())},default=str,ensure_ascii=False,indent=2).encode(),path,'json')
   continue
  im=image(path);ext='png'
  if op.startswith('to.'):
   ext=op.split('.')[1]
   with Image.open(path) as src:
    if getattr(src,'n_frames',1)>1 and ext in ['gif','webp','tiff']:
     frames=[ImageOps.exif_transpose(fr.copy()).convert('RGBA') for fr in ImageSequence.Iterator(src)]
     durations=[fr.info.get('duration',100) for fr in ImageSequence.Iterator(src)]
     output(image_bytes(frames[0],ext,save_all=True,append_images=frames[1:],duration=durations,**animation_loop(src,ext)),path,ext);continue
  elif op.startswith('rotate'):im=im.rotate(-int(op[6:]),expand=True)
  elif op=='flipH':im=ImageOps.mirror(im)
  elif op=='flipV':im=ImageOps.flip(im)
  elif op=='gray':
   alpha=im.convert('RGBA').getchannel('A');im=ImageOps.grayscale(im).convert('RGBA');im.putalpha(alpha)
  elif op in ['upscale2','upscale4']:
   factor=int(op[-1]);size=(im.width*factor,im.height*factor)
   if size[0]*size[1]>120_000_000:raise UserError('Risultato superiore a 120 megapixel.')
   im=im.resize(size,Image.Resampling.LANCZOS)
  elif op=='resize':
   w=int(number(p,'width',1920,1,20000));h=int(number(p,'height',0,0,20000)) or max(1,round(im.height*w/im.width))
   if w*h>120_000_000:raise UserError('Risultato troppo grande.')
   im=im.resize((w,h),Image.Resampling.LANCZOS)
  elif op in ['crop','redact']:
   x=int(number(p,'x',0,0,im.width-1));y=int(number(p,'y',0,0,im.height-1));w=int(number(p,'width',200,1,im.width-x));h=int(number(p,'height',200,1,im.height-y))
   if op=='crop':im=im.crop((x,y,x+w,y+h))
   else:im=im.convert('RGBA');ImageDraw.Draw(im).rectangle((x,y,x+w-1,y+h-1),fill='black')
  elif op=='edit':
   for key,cls in [('brightness',ImageEnhance.Brightness),('contrast',ImageEnhance.Contrast),('saturation',ImageEnhance.Color),('sharpness',ImageEnhance.Sharpness)]:im=cls(im).enhance(number(p,key,1,0,10))
  elif op=='annotate':
   im=im.convert('RGBA');ImageDraw.Draw(im).text((number(p,'x',20),number(p,'y',20)),nonempty(p,'text'),font=font(number(p,'size',32,1,1000)),fill=p.get('color','#ff8000'))
  elif op in ['watermark','tilemark']:
   im=im.convert('RGBA');layer=Image.new('RGBA',im.size);draw=ImageDraw.Draw(layer);fs=max(16,min(im.size)//15);text=nonempty(p,'text');fill=(255,255,255,int(number(p,'opacity',.35,0,1)*255))
   if op=='tilemark':
    for y in range(0,im.height,fs*5):
     for x in range(0,im.width,max(fs*5,int(draw.textlength(text,font=font(fs)))+fs)):draw.text((x,y),text,font=font(fs),fill=fill,stroke_width=1,stroke_fill=(0,0,0,fill[-1]))
   else:draw.text((im.width-20,im.height-20),text,font=font(fs),fill=fill,anchor='rd',stroke_width=1,stroke_fill=(0,0,0,fill[-1]))
   im=Image.alpha_composite(im,layer)
  elif op=='background':
   bg=Image.new('RGBA',im.size,p.get('color','#ffffff'));bg.alpha_composite(im.convert('RGBA'));im=bg
  elif op in ['removebg','blurbg']:
   mask=mask_subject(im);fg=im.convert('RGBA')
   if op=='removebg':fg.putalpha(ImageChops.multiply(fg.getchannel('A'),mask));im=fg
   else:im=Image.composite(fg,fg.filter(ImageFilter.GaussianBlur(20)),mask)
  elif op=='palette':
   colors=im.convert('RGB').resize((200,200)).quantize(colors=8).convert('RGB').getcolors(40000);colors=sorted(colors,reverse=True)
   output(json.dumps([{'hex':'#%02x%02x%02x'%rgb,'weight':count/40000} for count,rgb in colors],indent=2).encode(),path,'json');continue
  elif op=='metadata':
   output(json.dumps({'size':im.size,'mode':im.mode,'info':im.info,'exif':dict(im.getexif())},default=str,ensure_ascii=False,indent=2).encode(),path,'json');continue
  elif op=='editmeta':
   exif=im.getexif()
   for key,tag in [('artist',315),('copyright',33432),('description',270)]:exif[tag]=p.get(key,'')
   output(image_bytes(im,'jpg',exif=exif.tobytes()),path,'jpg');continue
  elif op=='strip':im.info.clear()
  elif op.startswith('quality'):
   output(image_bytes(im,'jpg',quality=int(op[7:])),path,'jpg');continue
  elif op=='compress':
   limit=int(number(p,'mb',1,.001,10000)*1_000_000)
   while True:
    best=None
    for q in [92,85,75,65,50,35,20,10]:
     data=image_bytes(im,'jpg',quality=q,optimize=True)
     if len(data)<=limit:best=data;break
    if best:break
    if im.size==(1,1):raise UserError('Limite troppo piccolo per questa immagine.')
    im=im.resize((max(1,int(im.width*.8)),max(1,int(im.height*.8))),Image.Resampling.LANCZOS)
   output(best,path,'jpg');continue
  else:raise UserError('Operazione immagine sconosciuta: '+op)
  output(image_bytes(im,ext),path,ext)

def process_vision(action,paths,p):
 op=action.split('.')[1]
 for path in paths:
  if op in ['ocr','qr']:
   result=vision(path,op);text='\n'.join(x['text'] for x in result) or ('Nessun codice QR o a barre riconosciuto.' if op=='qr' else 'Nessun testo riconosciuto.');output(text.encode(),path,'txt');continue
  im=image(path).convert('RGBA');draw=ImageDraw.Draw(im)
  for item in vision(path,'faces' if op=='faces' else 'ocr'):
   x,y,w,h=item['rect'];draw.rectangle((x*im.width,y*im.height,(x+w)*im.width,(y+h)*im.height),fill='black')
  output(image_bytes(im,'png'),path,'png')

def ffmpeg():
 import imageio_ffmpeg
 return imageio_ffmpeg.get_ffmpeg_exe()

def ff(args,timeout=3600):return run([ffmpeg(),'-nostdin','-hide_banner','-loglevel','error','-y']+args,timeout)

def duration(path):
 p=subprocess.run([ffmpeg(),'-hide_banner','-i',str(path)],capture_output=True)
 m=re.search(r'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)',p.stderr.decode(errors='replace'))
 if not m:raise UserError('Durata non leggibile.')
 h,m,s=map(float,m.groups());return h*3600+m*60+s

def vcodec(ext):
 return {'mp4':['-c:v','libx264','-preset','veryfast','-crf','20','-pix_fmt','yuv420p','-c:a','aac','-movflags','+faststart'],'mov':['-c:v','libx264','-preset','veryfast','-crf','20','-pix_fmt','yuv420p','-c:a','aac'],'mkv':['-c:v','libx264','-preset','veryfast','-crf','20','-c:a','aac'],'webm':['-c:v','libvpx-vp9','-deadline','realtime','-cpu-used','6','-crf','32','-b:v','0','-c:a','libopus'],'avi':['-c:v','mpeg4','-q:v','4','-c:a','libmp3lame'],'wmv':['-c:v','wmv2','-c:a','wmav2']}.get(ext,[])

def acodec(ext):return {'mp3':['-c:a','libmp3lame','-q:a','2'],'m4a':['-c:a','aac','-b:a','192k'],'wav':['-c:a','pcm_s16le'],'flac':['-c:a','flac'],'ogg':['-c:a','libvorbis','-q:a','5'],'opus':['-c:a','libopus','-b:a','128k'],'aiff':['-c:a','pcm_s16be'],'wma':['-c:a','wmav2','-b:a','192k']}[ext]

def tempo(speed):
 parts=[]
 while speed>2:parts.append('atempo=2');speed/=2
 while speed<.5:parts.append('atempo=0.5');speed/=.5
 parts.append(f'atempo={speed}');return ','.join(parts)

def transcription(path,ext,tmp):
 wav=Path(tmp)/'voice.wav';ff(['-i',path,'-vn','-ar','16000','-ac','1','-c:a','pcm_s16le',wav])
 cli=RES/('bin/arm64/whisper-cli' if platform.machine()=='arm64' else 'bin/whisper-cli');model=RES/'Models/ggml-tiny.bin';prefix=Path(tmp)/'transcript'
 if not cli.exists() or not model.exists():raise UserError('Motore trascrizione mancante.')
 run([cli,'-m',model,'-f',wav,'-l','auto','-t','4','-otxt','-osrt','-ovtt','-of',prefix],timeout=7200)
 return Path(str(prefix)+('.'+ext if ext else ''))

def trim_range(path,p):
 d=duration(path);start=number(p,'start',0,0);end=number(p,'end',5,start+.001)
 if start>=d or end>d+.02:raise UserError('Intervallo fuori dalla durata del file (%.3f s).'%d)
 return start,min(end,d)

def wav_codec(path):
 info=subprocess.run([ffmpeg(),'-hide_banner','-i',str(path)],capture_output=True).stderr.decode(errors='replace')
 line=next((line for line in info.splitlines() if 'Audio:' in line),'')
 if 'dbl' in line or 'pcm_f64' in line:return 'pcm_f64le'
 if 'flt' in line or 'pcm_f32' in line:return 'pcm_f32le'
 if '24 bit' in line or 'pcm_s24' in line:return 'pcm_s24le'
 if 's32' in line or 'pcm_s32' in line:return 'pcm_s32le'
 return 'pcm_s16le'

def process_media(action,paths,p):
 family,op=action.split('.',1)
 with tempfile.TemporaryDirectory(prefix='orbita-media-') as tmp:
  tmp=Path(tmp)
  if op=='join':
   # Re-encode to compatible intermediates, avoiding concat path injection.
   normalized=[]
   for i,path in enumerate(paths):
    dest=tmp/(str(i)+('.mp4' if family=='video' else '.wav'))
    if family=='video':
     info=subprocess.run([ffmpeg(),'-hide_banner','-i',str(path)],capture_output=True).stderr.decode(errors='replace')
     base=['-i',path]
     if 'Audio:' not in info:base+=['-f','lavfi','-i','anullsrc=r=48000:cl=stereo','-map','0:v:0','-map','1:a:0','-shortest']
     ff(base+['-vf','scale=1280:720:force_original_aspect_ratio=decrease,pad=1280:720:(ow-iw)/2:(oh-ih)/2,setsar=1','-r','30']+vcodec('mp4')+['-ar','48000','-ac','2',dest])
    else:ff(['-i',path,'-vn','-ar','48000','-ac','2',dest])
    normalized.append(dest)
   listing=tmp/'concat.txt';listing.write_text(''.join("file '"+x.name+"'\n" for x in normalized))
   ext='mp4' if family=='video' else 'wav';dest=tmp/('joined.'+ext)
   ff(['-f','concat','-safe','1','-i',listing,'-c','copy',dest]);publish(dest,paths[0],ext);return
  for path in paths:
   if family=='transcribe':publish(transcription(path,op,tmp),path,op);continue
   if op=='autosubs':
    sub=transcription(path,'srt',tmp);dest=tmp/'subtitled.mp4'
    ff(['-i',path,'-vf',f"subtitles=filename='{sub}'"]+vcodec('mp4')+[dest]);publish(dest,path,'mp4');continue
   ext=op.split('.')[1] if op.startswith('to.') else ('mp4' if family=='video' else 'm4a')
   if op in ['strip','tags']:ext=Path(path).suffix.lower().lstrip('.')
   dest=tmp/('result.'+ext);args=['-i',path];opts=[]
   if family=='audio':
    opts=['-vn']+(acodec(ext) if op not in ['strip','tags'] else ['-c:a','copy'])
    if op=='to.wav':opts=['-vn','-c:a',wav_codec(path)]
    if op=='trim':
     start,end=trim_range(path,p);opts+=['-ss',str(start),'-t',str(end-start)]
    elif op=='bleep':
     start=number(p,'start',0,0);end=number(p,'end',5,start+.001)
     args+=['-f','lavfi','-i',f'sine=frequency=1000:sample_rate=48000:duration={end-start}'];opts+=['-filter_complex',f'[0:a]volume=0:enable=\'between(t,{start},{end})\'[a];[1:a]adelay={int(start*1000)}|{int(start*1000)}[b];[a][b]amix=inputs=2:duration=first:normalize=0[out]','-map','[out]']
    elif op=='fade':
     d=duration(path);t=number(p,'fade',1,.01,d/2);opts+=['-af',f'afade=t=in:d={t},afade=t=out:st={d-t}:d={t}']
    elif op=='normalize':opts+=['-af','loudnorm=I=-16:TP=-1.5:LRA=11']
    elif op=='denoise':opts+=['-af','afftdn=nf=-25']
    elif op=='speed':opts+=['-af',tempo(number(p,'speed',1.5,.1,10))]
    elif op=='pitch':
     ratio=2**(number(p,'semitones',2,-24,24)/12);opts+=['-af',f'aresample=48000,asetrate={48000*ratio},aresample=48000,'+tempo(1/ratio)]
    elif op in ['mono','stereo']:opts+=['-ac','1' if op=='mono' else '2']
    elif op=='compress':opts=['-vn','-c:a','aac','-b:a','128k']
    elif op=='silence':
     threshold=number(p,'db',-45,-90,-1);opts+=['-af',f'silenceremove=start_periods=1:start_duration=0.05:start_threshold={threshold}dB,areverse,silenceremove=start_periods=1:start_duration=0.05:start_threshold={threshold}dB,areverse']
    elif op=='volume':opts+=['-af',f"volume={number(p,'volume',1,0,10)}",'-ac',str(int(number(p,'channels',2,1,2)))]
    elif op=='waveform':
     ext='mp4';dest=tmp/'waveform.mp4';opts=['-filter_complex','[0:a]aformat=channel_layouts=mono,showwaves=s=1280x720:mode=line:colors=0x9999ff:rate=30[v]','-map','[v]','-map','0:a']+vcodec('mp4')+['-shortest']
    elif op=='tags':
     opts=['-map','0','-c','copy']
     for key in ['title','artist','album']:opts+=['-metadata',key+'='+str(p.get(key,''))]
    elif op=='artwork':
     args+=['-i',nonempty(p,'file')];opts=acodec('mp3')+['-map','0:a','-map','1:v','-c:v','mjpeg','-id3v2_version','3','-disposition:v','attached_pic'];ext='mp3';dest=tmp/'result.mp3'
    elif op=='strip':opts+=['-map_metadata','-1']
    elif not op.startswith('to.'):raise UserError('Operazione audio sconosciuta.')
   else:
    opts=vcodec(ext)
    if op.startswith('to.') and ext in ['mp4','mov','mkv','webm','avi','wmv']:opts+=['-vf','pad=ceil(iw/2)*2:ceil(ih/2)*2']
    if ext=='gif':opts=['-vf','fps=12,scale=640:-1:flags=lanczos,split[s0][s1];[s0]palettegen[p];[s1][p]paletteuse','-loop','0']
    elif op=='trim':
     start,end=trim_range(path,p);opts+=['-ss',str(start),'-t',str(end-start)]
    elif op=='split':
     point=number(p,'start',5,.001,duration(path)-.001)
     for i,limits in enumerate([['-t',str(point)],['-ss',str(point)]],1):
      part=tmp/f'part{i}.mp4';ff(args+limits+vcodec('mp4')+[part]);publish(part,path,'mp4',f'-parte-{i}')
     continue
    elif op in ['crop','redact']:
     x=int(number(p,'x',0,0));y=int(number(p,'y',0,0));w=int(number(p,'width',200,1));h=int(number(p,'height',200,1))
     if op=='crop':
      frame=tmp/'dimensions.png';ff(['-i',path,'-map','0:v:0','-frames:v','1',frame])
      with Image.open(frame) as im:iw,ih=im.size
      if x+w>iw or y+h>ih:raise UserError('Il ritaglio supera i bordi del video.')
      vf=f'crop={w}:{h}:{x}:{y}:exact=1,pad=ceil(iw/2)*2:ceil(ih/2)*2'
     else:
      start=number(p,'start',0,0);end=number(p,'end',5,start+.001);vf=f"drawbox=x={x}:y={y}:w={w}:h={h}:color=black:t=fill:enable='between(t,{start},{end})'"
     opts+=['-vf',vf]
    elif op=='frames':
     fps=number(p,'fps',1,.01,60)
     if duration(path)*fps>2000:raise UserError('Massimo 2000 fotogrammi per operazione; riduci la frequenza.')
     destdir=folder(path,'-fotogrammi');ff(args+['-vf',f'fps={fps}',destdir/'frame-%05d.png']);continue
    elif op=='compress':
     limit=number(p,'mb',8,.05,100000)*1_000_000;d=duration(path);bitrate=int(limit*8/d*.92)-64000
     if bitrate<20000:raise UserError('Limite troppo piccolo per la durata del video.')
     for attempt in range(4):
      opts=['-c:v','libx264','-preset','veryfast','-pix_fmt','yuv420p','-b:v',str(bitrate),'-c:a','aac','-b:a','64k','-movflags','+faststart']
      ff(args+opts+[dest])
      if dest.stat().st_size<=limit:break
      bitrate=int(bitrate*limit/dest.stat().st_size*.9)
     else:raise UserError('Impossibile rispettare il limite richiesto.')
     publish(dest,path,'mp4');continue
    elif op=='speed':
     speed=number(p,'speed',2,.1,10)
     # Audio is optional; use independent filters so silent videos also work.
     opts+=['-vf',f'setpts=PTS/{speed}','-af',tempo(speed)]
    elif op=='rotate':opts+=['-vf','transpose=1']
    elif op=='mute':opts+=['-an']
    elif op=='strip':opts=['-map','0','-c','copy','-map_metadata','-1']
    elif op in ['addsubs','burnsubs']:
     sub=tmp/('subtitles'+Path(nonempty(p,'file')).suffix);shutil.copyfile(p['file'],sub)
     if op=='burnsubs':opts+=['-vf',f"subtitles=filename='{sub}'"]
     else:args+=['-i',sub];opts+=['-map','0:v','-map','0:a?','-map','0:s?','-map','1:0','-c:s','mov_text']
    elif op=='extractsubs':ext='srt';dest=tmp/'subtitles.srt';opts=['-map','0:s:0','-c:s','srt']
    elif not op.startswith('to.'):raise UserError('Operazione video sconosciuta.')
   ff(args+opts+[dest]);publish(dest,path,ext)

def pdf_text(doc):return '\n\n'.join(page.get_text(sort=True) for page in doc)

def text_pdf(text):
 import pymupdf as fitz
 doc=fitz.open();story=fitz.Story(html='<html><body>'+''.join('<p>'+html.escape(x)+'</p>' for x in text.split('\n'))+'</body></html>');buf=io.BytesIO();writer=fitz.DocumentWriter(buf)
 more=True
 while more:
  dev=writer.begin_page(fitz.Rect(0,0,595,842));more,_=story.place(fitz.Rect(40,40,555,802));story.draw(dev);writer.end_page()
 writer.close();return buf.getvalue()

def insert_unicode(page,point,text,size=14,**kwargs):
 import pymupdf as fitz
 font=fitz.Font('cjk')
 if any(not font.has_glyph(ord(c)) for c in text if not c.isspace()):raise UserError('Il testo contiene caratteri non disponibili nel font PDF.')
 x,y=point;lines=text.splitlines() or [''];leading=size*1.2
 box=fitz.Rect(x,y-size*font.ascender,x+max(font.text_length(line,fontsize=size) for line in lines),y+(len(lines)-1)*leading-size*font.descender)
 if box.is_empty or not page.rect.contains(box):raise UserError('Il testo non entra nella pagina: modifica posizione o dimensione.')
 page.insert_font(fontname='OrbitaUnicode',fontbuffer=font.buffer)
 page.insert_text(point,text,fontname='OrbitaUnicode',fontsize=size,**kwargs)

def process_pdf(action,paths,p):
 import pymupdf as fitz
 op=action.split('.',1)[1]
 if op=='merge':
  out=fitz.open();toc=[]
  for path in paths:
   with fitz.open(path) as doc:
    if doc.needs_pass:raise UserError('Sblocca prima i PDF protetti.')
    offset=len(out);out.insert_pdf(doc)
    toc.extend([[level,title,page+offset if page>0 else page] for level,title,page in doc.get_toc()])
  if toc:out.set_toc(toc)
  output(out.tobytes(garbage=4,deflate=True),paths[0],'pdf');return
 for path in paths:
  with fitz.open(path) as doc:
   if doc.needs_pass and not doc.authenticate(str(p.get('password',''))):raise UserError('Password PDF necessaria o errata.')
   if op.startswith('to.'):
    ext=op.split('.')[1];text=pdf_text(doc)
    if ext in ['png','jpg']:
     for i,page in enumerate(doc):output(page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).tobytes(ext),path,ext,f'-pagina-{i+1}')
    elif ext in ['txt','md']:output(text.encode(),path,ext)
    elif ext=='docx':
     from docx import Document
     d=Document()
     for page in doc:
      for line in page.get_text(sort=True).splitlines():d.add_paragraph(line)
      if page.number<doc.page_count-1:d.add_page_break()
     b=io.BytesIO();d.save(b);output(b.getvalue(),path,ext)
    elif ext=='xlsx':
     from openpyxl import Workbook
     book=Workbook();book.remove(book.active)
     for page in doc:
      tables=page.find_tables().tables
      for i,table in enumerate(tables):
       ws=book.create_sheet(f'P{page.number+1} T{i+1}')
       for row in table.extract():append_sheet_row(ws,row)
     if not book.worksheets:raise UserError('Nessuna tabella riconosciuta nel PDF. Se è una scansione, esegui prima OCR.')
     b=io.BytesIO();book.save(b);output(b.getvalue(),path,ext)
    elif ext=='pptx':
     from pptx import Presentation
     from pptx.util import Inches
     prs=Presentation();first=doc[0].rect;prs.slide_width=Inches(first.width/72);prs.slide_height=Inches(first.height/72)
     for page in doc:
      slide=prs.slides.add_slide(prs.slide_layouts[6]);pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5));scale=min(prs.slide_width/page.rect.width,prs.slide_height/page.rect.height);w=int(page.rect.width*scale);h=int(page.rect.height*scale);slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')),int((prs.slide_width-w)/2),int((prs.slide_height-h)/2),width=w,height=h)
     b=io.BytesIO();prs.save(b);output(b.getvalue(),path,ext)
    elif ext=='epub':output(make_epub(text,Path(path).stem),path,ext)
    continue
   if op=='split':
    for i in range(len(doc)):
     new=fitz.open();new.insert_pdf(doc,from_page=i,to_page=i);output(new.tobytes(),path,'pdf',f'-pagina-{i+1}')
    continue
   if op=='select':
    pages=[]
    for token in nonempty(p,'pages').split(','):
     if '-' in token:
      a,b=map(int,token.split('-'));pages.extend(range(a-1,b))
     else:pages.append(int(token)-1)
    if not pages or min(pages)<0 or max(pages)>=len(doc):raise UserError('Numeri di pagina fuori intervallo.')
    doc.select(pages)
   elif op=='rotate':
    for page in doc:page.set_rotation((page.rotation+90)%360)
   elif op=='reverse':doc.select(list(reversed(range(len(doc)))))
   elif op in ['numbers','watermark']:
    for i,page in enumerate(doc):
     if op=='numbers':page.insert_text((page.rect.width/2, page.rect.height-20),str(i+int(number(p,'start',1,0))),fontsize=10)
     else:insert_unicode(page,(40,page.rect.height/2),nonempty(p,'text'),40,color=(.6,.6,.6),fill_opacity=.3)
   elif op in ['compress','gray']:
    new=fitz.open();dpi=number(p,'dpi',120,36,600) if op=='compress' else 144
    for page in doc:
     pix=page.get_pixmap(dpi=int(dpi),colorspace=fitz.csGRAY if op=='gray' else fitz.csRGB)
     np=new.new_page(width=page.rect.width,height=page.rect.height);np.insert_image(np.rect,stream=pix.tobytes('jpeg',jpg_quality=75))
    output(new.tobytes(garbage=4,deflate=True),path,'pdf');continue
   elif op=='images':
    seen=set()
    for page in doc:
     for info in page.get_images():
      xref=info[0]
      if xref in seen:continue
      seen.add(xref);img=doc.extract_image(xref)
      if info[1]:
       pix=fitz.Pixmap(doc,xref);mask=fitz.Pixmap(doc,info[1]);pix=fitz.Pixmap(pix,mask);output(pix.tobytes('png'),path,'png',f'-immagine-{len(seen)}')
      else:output(img['image'],path,img['ext'],f'-immagine-{len(seen)}')
    if not seen:raise UserError('Il PDF non contiene immagini raster estraibili.')
    continue
   elif op in ['metadata','forms']:
    data=doc.metadata if op=='metadata' else [{'page':page.number+1,'name':w.field_name,'value':w.field_value,'type':w.field_type_string} for page in doc for w in (page.widgets() or [])]
    output(json.dumps(data,ensure_ascii=False,indent=2).encode(),path,'json');continue
   elif op=='strip':doc.set_metadata({});doc.del_xml_metadata()
   elif op=='editmeta':doc.set_metadata({**doc.metadata,**{k:str(p.get(k,'')) for k in ['title','author','subject']}})
   elif op=='redact':
    found=0
    for page in doc:
     for rect in page.search_for(nonempty(p,'text')):page.add_redact_annot(rect,fill=(0,0,0));found+=1
     page.apply_redactions()
    if not found:raise UserError('Testo non trovato: per le scansioni esegui prima OCR.')
   elif op in ['addtext','sign']:
    index=int(number(p,'page',1,1,len(doc)))-1;page=doc[index];x=number(p,'x',50,0);y=number(p,'y',50,0)
    if op=='addtext':insert_unicode(page,(x,y),nonempty(p,'text'),number(p,'size',14,1,500))
    else:
     box=fitz.Rect(x,y,x+number(p,'width',200,1),y+number(p,'height',80,1))
     if not page.rect.contains(box):raise UserError('La firma non entra nella pagina.')
     page.insert_image(box,filename=nonempty(p,'file'),keep_proportion=True)
   elif op=='fill':
    values=json.loads(nonempty(p,'values'));found=0
    if not isinstance(values,dict):raise UserError('Inserisci un oggetto JSON con nome campo e valore.')
    for page in doc:
     for widget in page.widgets() or []:
      if widget.field_name in values:widget.field_value=values[widget.field_name];widget.update();found+=1
    if not found:raise UserError('Nessun campo corrispondente. Usa «Elenca campi modulo».')
   elif op=='ocr':
    with tempfile.TemporaryDirectory(prefix='orbita-ocr-') as tmp:
     for page in doc:
      png=Path(tmp)/'page.png';page.get_pixmap(dpi=144).save(png)
      for item in vision(png):
       text=item['text'];x,y,w,h=item['rect'];rect=fitz.Rect(x*page.rect.width,y*page.rect.height,(x+w)*page.rect.width,(y+h)*page.rect.height)
       if any((r & rect).get_area()>0 for r in page.search_for(text)):continue
       font=fitz.Font('cjk');fs=min(30,rect.height*.7,rect.width/max(.001,font.text_length(text)))
       if fs<=0:continue
       insert_unicode(page,(rect.x0,rect.y0+fs*font.ascender),text,fs,render_mode=3)
   elif op not in ['unlock','password']:raise UserError('Operazione PDF sconosciuta.')
   kwargs=dict(garbage=4,deflate=True)
   if op=='password':
    password=nonempty(p,'password');kwargs.update(encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw=password,user_pw=password,permissions=fitz.PDF_PERM_PRINT|fitz.PDF_PERM_COPY)
   else:kwargs['encryption']=fitz.PDF_ENCRYPT_NONE
   output(doc.tobytes(**kwargs),path,'pdf')

def append_sheet_row(sheet,values):
 from openpyxl.cell import Cell
 cells=[]
 for value in values:
  if isinstance(value,str) and len(value)>32767:raise UserError('Una cella supera i 32767 caratteri consentiti da XLSX. Usa CSV o JSON per conservarla interamente.')
  cell=Cell(sheet,value=value)
  if isinstance(value,str):cell.data_type='s'
  cells.append(cell)
 sheet.append(cells)

def make_epub(text,title):
 from ebooklib import epub
 book=epub.EpubBook();book.set_identifier(hashlib.sha256(text.encode()).hexdigest());book.set_title(title);book.set_language('it')
 chapter=epub.EpubHtml(title=title,file_name='content.xhtml',lang='it');chapter.content='<h1>'+html.escape(title)+'</h1>'+''.join('<p>'+html.escape(line)+'</p>' for line in text.splitlines());book.add_item(chapter);book.toc=(chapter,);book.spine=['nav',chapter];book.add_item(epub.EpubNcx());book.add_item(epub.EpubNav())
 with tempfile.TemporaryDirectory() as tmp:
  file=Path(tmp)/'book.epub';epub.write_epub(file,book);return file.read_bytes()

def libreoffice():
 choices=[RES/('Office-arm64/LibreOffice.app/Contents/MacOS/soffice' if platform.machine()=='arm64' else 'Office/LibreOffice.app/Contents/MacOS/soffice'),Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')]
 for path in choices:
  if path.exists():return path
 raise UserError('Motore LibreOffice non presente nel pacchetto.')

def office_convert(path,ext,tmp):
 outdir=Path(tmp)/'office';outdir.mkdir(exist_ok=True)
 # Isolated profile prevents interference with other LibreOffice processes.
 profile=(Path(tmp)/'lo-profile').as_uri()
 run([libreoffice(),'-env:UserInstallation='+profile,'--headless','--convert-to',ext,'--outdir',outdir,path],timeout=600)
 result=outdir/(Path(path).stem+('.'+ext if ext else ''))
 if not result.exists():raise UserError('Questo documento non è stato convertito da LibreOffice.')
 return result

def read_text_file(source):
 data=Path(source).read_bytes()
 encoding='utf-32' if data.startswith((b'\xff\xfe\x00\x00',b'\x00\x00\xfe\xff')) else 'utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig'
 return data.decode(encoding)

def embed_html_images(source):
 import base64,mimetypes
 from urllib.parse import unquote,urlparse
 text=Path(source).read_text(encoding='utf-8')
 def replace(m):
  value=html.unescape(m[2]);parsed=urlparse(value)
  if parsed.scheme in ['data','http','https']:return m[0]
  path=Path(unquote(parsed.path))
  if not path.is_absolute():path=Path(source).parent/path
  if not path.is_file():raise UserError('Immagine del documento mancante: '+value)
  mime=mimetypes.guess_type(str(path))[0] or 'application/octet-stream'
  return m[1]+'data:'+mime+';base64,'+base64.b64encode(path.read_bytes()).decode()+m[3]
 return re.sub(r"(?i)(\bsrc\s*=\s*['\"])([^'\"]+)(['\"])",replace,text).encode()

def process_document(action,paths,p):
 family,_,ext=action.split('.')
 import pymupdf as fitz
 import pypandoc
 for path in paths:
  with tempfile.TemporaryDirectory(prefix='orbita-doc-') as tmp:
   source=Path(path);source_ext=source.suffix.lower().lstrip('.')
   if family=='slides':
    if ext in ['png','jpg']:
     pdf=office_convert(source,'pdf',tmp)
     with fitz.open(pdf) as doc:
      for page in doc:output(page.get_pixmap(dpi=144).tobytes(ext),path,ext,f'-slide-{page.number+1}')
    else:publish(office_convert(source,ext,tmp),path,ext)
    continue
   if source_ext in ['doc','docx','odt','rtf','pages'] and ext in ['pdf','docx','odt','rtf','html','txt']:
    converted=office_convert(source,ext,tmp)
    if ext=='html':output(embed_html_images(converted),path,'html')
    else:publish(converted,path,ext)
    continue
   if source_ext in ['rtf','doc','pages']:
    source=office_convert(source,'docx',tmp);source_ext='docx'
   if source_ext=='txt':
    # Plain text must preserve markup characters and whitespace literally.
    literal=read_text_file(source)
    source=Path(tmp)/'literal.html';source.write_text('<html><body><pre>'+html.escape(literal)+'</pre></body></html>');source_ext='html'
   fmt={'md':'markdown','htm':'html'}.get(source_ext,source_ext)
   if ext=='pdf':
    h=pypandoc.convert_file(str(source),'html',format=fmt,extra_args=['--standalone','--embed-resources','--resource-path='+str(source.parent)])
    story=fitz.Story(html=h);buf=io.BytesIO();writer=fitz.DocumentWriter(buf);more=True
    while more:
     dev=writer.begin_page(fitz.Rect(0,0,595,842));more,_=story.place(fitz.Rect(40,40,555,802));story.draw(dev);writer.end_page()
    writer.close();output(buf.getvalue(),path,'pdf')
   else:
    dest=Path(tmp)/('result.'+ext)
    pypandoc.convert_file(str(source),{'md':'markdown','txt':'plain'}.get(ext,ext),format=fmt,outputfile=str(dest),extra_args=['--standalone'])
    publish(dest,path,ext)

def read_sheet(path):
 ext=Path(path).suffix.lower()
 if ext in ['.csv','.tsv']:
  with open(path,encoding='utf-8-sig',newline='') as f:return list(csv.reader(f,delimiter='\t' if ext=='.tsv' else ','))
 if ext=='.json':
  data=json.loads(Path(path).read_text())
  if not isinstance(data,list):raise UserError('Il JSON deve contenere una lista di record o righe.')
  if data and isinstance(data[0],dict):
   keys=list(dict.fromkeys(k for row in data for k in row));return [keys]+[[row.get(k,'') for k in keys] for row in data]
  if not all(isinstance(x,list) for x in data):raise UserError('Il JSON deve contenere righe o oggetti.')
  return data
 if ext=='.xlsx':
  from openpyxl import load_workbook
  formulas=load_workbook(path,read_only=True,data_only=False)
  book=load_workbook(path,read_only=True,data_only=True)
  try:
   sheet=book.worksheets[0];rows=[list(r) for r in sheet.values]
   missing=any(c.data_type=='f' and sheet.cell(c.row,c.column).value is None for row in formulas.worksheets[0] for c in row)
  finally:formulas.close();book.close()
  if missing:
   with tempfile.TemporaryDirectory(prefix='orbita-recalc-') as tmp:
    recalculated=office_convert(path,'xlsx',tmp);book=load_workbook(recalculated,read_only=True,data_only=True)
    try:
     sheet=book.worksheets[0];rows=[list(r) for r in sheet.values]
     if any(c.data_type=='e' for row in sheet for c in row):raise UserError('Il foglio contiene formule non calcolabili. Aprilo e correggile prima di esportare.')
    finally:book.close()
  return rows
 if ext=='.ods':
  from odf.opendocument import load
  from odf.table import Table,TableRow
  from odf.namespaces import TABLENS
  from odf import teletype
  table=load(path).spreadsheet.getElementsByType(Table)[0];rows=[];cells=0
  for row in table.getElementsByType(TableRow):
   values=[];repeat=int(row.getAttribute('numberrowsrepeated') or 1)
   for cell in row.childNodes:
    if getattr(cell,'qname',None) not in [(TABLENS,'table-cell'),(TABLENS,'covered-table-cell')]:continue
    count=int(cell.getAttribute('numbercolumnsrepeated') or 1)
    if count<1 or repeat<1 or (len(values)+count)*repeat+cells>2_000_000:raise UserError('Foglio troppo grande: oltre 2 milioni di celle. Nessun dato è stato troncato.')
    kind=cell.getAttribute('valuetype') if cell.qname[1]=='table-cell' else None
    value=teletype.extractText(cell)
    if kind in ['float','percentage','currency'] and cell.getAttribute('value') is not None:value=float(cell.getAttribute('value'))
    elif kind=='boolean':value=cell.getAttribute('booleanvalue')=='true'
    elif not value and kind=='date':value=cell.getAttribute('datevalue') or ''
    elif not value and kind=='time':value=cell.getAttribute('timevalue') or ''
    elif not value and kind=='string':value=cell.getAttribute('stringvalue') or ''
    values.extend([value]*count)
   cells+=len(values)*repeat
   if len(rows)+repeat>1_048_576:raise UserError('Foglio oltre un milione di righe. Nessun dato è stato troncato.')
   rows.extend([values.copy() for _ in range(repeat)])
  return rows
 raise UserError('Formato tabella non supportato.')

def process_sheet(action,paths,p):
 ext=action.split('.')[-1]
 for path in paths:
  with tempfile.TemporaryDirectory(prefix='orbita-sheet-') as tmp:
   if ext in ['pdf','ods','xlsx'] and Path(path).suffix.lower() in ['.xlsx','.xls','.ods','.numbers']:
    publish(office_convert(path,ext,tmp),path,ext);continue
   read_path=office_convert(path,'xlsx',tmp) if Path(path).suffix.lower() in ['.xls','.numbers'] else path
   rows=read_sheet(read_path)
   if ext in ['csv','tsv']:
    text=io.StringIO();writer=csv.writer(text,delimiter='\t' if ext=='tsv' else ',');writer.writerows(rows);output(text.getvalue().encode('utf-8-sig'),path,ext)
   elif ext=='json':output(json.dumps(rows,ensure_ascii=False,default=str,indent=2).encode(),path,ext)
   elif ext=='ods':
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table,TableRow,TableCell
    from odf.text import P
    doc=OpenDocumentSpreadsheet();table=Table(name='Dati')
    for row in rows:
     r=TableRow()
     for value in row:c=TableCell(valuetype='string');c.addElement(P(text='' if value is None else str(value)));r.addElement(c)
     table.addElement(r)
    doc.spreadsheet.addElement(table);dest=Path(tmp)/'sheet.ods';doc.save(str(dest));publish(dest,path,ext)
   elif ext=='xlsx':
    from openpyxl import Workbook
    book=Workbook()
    for row in rows:append_sheet_row(book.active,row)
    b=io.BytesIO();book.save(b);output(b.getvalue(),path,ext)
   elif ext=='pdf':
    from openpyxl import Workbook
    book=Workbook()
    for row in rows:append_sheet_row(book.active,row)
    source=Path(tmp)/'sheet.xlsx';book.save(source);publish(office_convert(source,'pdf',tmp),path,ext)

def process_subtitle(action,paths,p):
 stamp=re.compile(r'(?:(\d{2,}):)?(\d{2}):(\d{2})[,.](\d{3})')
 for path in paths:
  text=Path(path).read_text(encoding='utf-8-sig');ext=action.split('.')[-1]
  if ext=='shift':
   offset=number(p,'seconds',1,-100000,100000)
   def change(m):
    h,mi,s,ms=m.groups();total=max(0,round((int(h or 0)*3600+int(mi)*60+int(s)+int(ms)/1000+offset)*1000))
    h,total=divmod(total,3600000);mi,total=divmod(total,60000);sec,ms=divmod(total,1000);sep=',' if ',' in m[0] else '.'
    return f'{h:02}:{mi:02}:{sec:02}{sep}{ms:03}'
   lines=[]
   for line in text.splitlines(keepends=True):
    if re.match(r'^\s*(?:\d{2,}:)?\d{2}:\d{2}[,.]\d{3}\s+-->',line):line=stamp.sub(change,line,count=2)
    lines.append(line)
   text=''.join(lines);ext=Path(path).suffix[1:]
  else:
   with tempfile.TemporaryDirectory(prefix='orbita-subs-') as tmp:
    dest=Path(tmp)/('result.'+ext);ff(['-i',path,dest]);text=dest.read_text()
  output(text.encode(),path,ext)

def archive_members(paths):
 for path in paths:
  path=Path(path)
  if path.is_symlink():raise UserError('I link simbolici non sono inclusi negli archivi.')
  if path.is_dir():
   yield path,path.name+'/'
   for child in path.rglob('*'):
    if child.is_symlink():raise UserError('La cartella contiene link simbolici: '+str(child))
    if child.is_file() or child.is_dir():yield child,str(Path(path.name)/child.relative_to(path))+('/' if child.is_dir() else '')
  else:yield path,path.name

def safe_archive_name(name):
 p=Path(name)
 if p.is_absolute() or '..' in p.parts or '\\' in name or '\x00' in name:raise UserError('Archivio con percorsi non sicuri.')

def process_archive(action,paths,p):
 op=action.split('.',1)[1]
 with tempfile.TemporaryDirectory(prefix='orbita-archive-') as tmp:
  tmp=Path(tmp)
  if op=='extract':
   for path in paths:
    dest=folder(path,'-estratto');suffix=Path(path).suffix.lower()
    try:
     if zipfile.is_zipfile(path):
      import pyzipper
      with pyzipper.AESZipFile(path) as z:
       if p.get('password'):z.setpassword(str(p['password']).encode())
       infos=z.infolist()
       if sum(i.file_size for i in infos)>20_000_000_000:raise UserError('Archivio espanso oltre 20 GB.')
       for i in infos:
        safe_archive_name(i.filename)
        if (i.external_attr>>16)&0o170000==0o120000:raise UserError('Link simbolico nell’archivio.')
       z.extractall(dest)
     elif tarfile.is_tarfile(path):
      with tarfile.open(path) as t:
       if sum(i.size for i in t.getmembers())>20_000_000_000:raise UserError('Archivio espanso oltre 20 GB.')
       for i in t.getmembers():
        safe_archive_name(i.name)
        if not (i.isfile() or i.isdir()):raise UserError('Tipo speciale non consentito nell’archivio.')
       t.extractall(dest,filter='data')
     elif suffix=='.7z':
      import py7zr
      with py7zr.SevenZipFile(path,password=p.get('password') or None) as z:
       for name in z.getnames():safe_archive_name(name)
       if sum(getattr(item,'uncompressed',0) or 0 for item in z.list())>20_000_000_000:raise UserError('Archivio espanso oltre 20 GB.')
       for item in z.files:
        if item.is_symlink or getattr(item,'is_junction',False):raise UserError('Link non consentito nell’archivio.')
       for item in z.list():
        if getattr(item,'is_symlink',False) or getattr(item,'uncompressed',0)>20_000_000_000:raise UserError('Elemento archivio non consentito.')
       z.extractall(dest)
     elif suffix=='.rar':
      # macOS libarchive handles RAR without shell interpolation.
      names=run(['/usr/bin/tar','-tf',path]).decode().splitlines()
      for name in names:safe_archive_name(name)
      listing=run(['/usr/bin/tar','-tvf',path]).decode().splitlines()
      if any(line and line[0] not in ['-','d'] for line in listing):raise UserError('Archivio con link o file speciali non consentito.')
      run(['/usr/bin/tar','-xf',path,'-C',dest])
     elif suffix in ['.gz','.bz2','.xz']:
      opener={'.gz':gzip.open,'.bz2':bz2.open,'.xz':lzma.open}[suffix]
      with opener(path,'rb') as src, (dest/Path(path).stem).open('xb') as out:shutil.copyfileobj(src,out)
     else:raise UserError('Archivio non riconosciuto.')
    except Exception:
     # Keep partial extraction visible, never touch pre-existing paths.
     raise
   return
  members=list(archive_members(paths));names=[name for _,name in members]
  if len(names)!=len(set(names)):raise UserError('File con nomi uguali: comprimi separatamente o rinomina prima.')
  ext='zip' if op=='password' else op;dest=tmp/('archive.'+ext)
  if op=='password':
   import pyzipper
   with pyzipper.AESZipFile(dest,'w',compression=pyzipper.ZIP_DEFLATED,encryption=pyzipper.WZ_AES) as z:
    z.setpassword(nonempty(p,'password').encode())
    for file,name in members:z.write(file,name)
  elif op=='zip':
   with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
    for file,name in members:z.write(file,name)
  else:
   with tarfile.open(dest,'w:gz' if op=='tar.gz' else 'w') as t:
    for file,name in members:t.add(file,arcname=name,recursive=False)
  publish(dest,paths[0],ext)

def execute(req):
 global DESTINATIONS
 DESTINATIONS={str(Path(k).resolve()):v for k,v in req.get('destinations',{}).items()}
 from catalog import A
 ids={a['id'] for a in A};action=req['action'];paths=req.get('paths',[]);p=req.get('params',{})
 if action not in ids:raise UserError('Strumento non riconosciuto.')
 if not paths or len(paths)>1000:raise UserError('Seleziona da 1 a 1000 file.')
 # Reject links before resolution for operations that change names or traverse archives.
 if action=='files.renameOriginals' or action.startswith('archive.'):
  if any(Path(x).is_symlink() for x in paths):raise UserError('Seleziona il file originale, non un collegamento simbolico.')
 paths=[str(Path(x).resolve(strict=True)) for x in paths]
 if action.startswith('image.'):process_image(action,paths,p)
 elif action.startswith('vision.'):process_vision(action,paths,p)
 elif action.startswith(('video.','audio.','transcribe.')):process_media(action,paths,p)
 elif action.startswith('pdf.'):process_pdf(action,paths,p)
 elif action.startswith(('document.','slides.')):process_document(action,paths,p)
 elif action.startswith('sheet.'):process_sheet(action,paths,p)
 elif action.startswith('subtitle.'):process_subtitle(action,paths,p)
 elif action.startswith('archive.'):process_archive(action,paths,p)
 elif action=='files.strip':
  for path in paths:
   ext=Path(path).suffix.lower()
   if ext=='.pdf':process_pdf('pdf.strip',[path],p)
   elif ext in ['.png','.jpg','.jpeg','.tif','.tiff','.gif','.webp','.heic','.avif','.bmp']:process_image('image.strip',[path],p)
   elif ext in ['.mp4','.mov','.mkv','.webm','.avi','.wmv']:process_media('video.strip',[path],p)
   elif ext in ['.mp3','.m4a','.wav','.flac','.ogg','.opus','.aiff','.wma']:process_media('audio.strip',[path],p)
   elif ext in ['.docx','.xlsx','.pptx','.odt','.ods','.odp']:
    b=io.BytesIO()
    with zipfile.ZipFile(path) as src,zipfile.ZipFile(b,'w',zipfile.ZIP_DEFLATED) as out:
     for info in src.infolist():
      data=src.read(info.filename)
      if info.filename=='docProps/custom.xml':continue
      if info.filename in ['[Content_Types].xml','_rels/.rels']:
       from xml.etree import ElementTree as ET
       tree=ET.fromstring(data)
       for node in list(tree):
        if node.get('PartName')=='/docProps/custom.xml' or node.get('Type','').endswith('/custom-properties'):tree.remove(node)
       data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
      if info.filename=='docProps/core.xml':data=b'<?xml version="1.0" encoding="UTF-8"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"/>'
      if info.filename=='meta.xml':data=b'<?xml version="1.0" encoding="UTF-8"?><office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"><office:meta/></office:document-meta>'
      out.writestr(info,data)
    output(b.getvalue(),path,ext[1:])
   else:raise UserError('Rimozione metadati non disponibile per questo formato. Nessuna copia modificata.')
 elif action=='files.hash':
  for path in paths:
   h=hashlib.sha256()
   with open(path,'rb') as file:
    for chunk in iter(lambda:file.read(1024*1024),b''):h.update(chunk)
   output((h.hexdigest()+'  '+Path(path).name+'\n').encode(),path,'sha256')
 elif action=='files.renameOriginals':
  prefix=nonempty(p,'prefix')
  if '/' in prefix or '\\' in prefix or prefix in ['.','..']:raise UserError('Prefisso non valido.')
  for i,path in enumerate(paths,int(number(p,'start',1,0))):
   source=Path(path)
   if not source.is_file():raise UserError('La rinomina in serie accetta file, non cartelle.')
   for collision in range(10000):
    target=source.with_name(prefix+f'-{i:04}'+(f'-{collision}' if collision else '')+source.suffix)
    if target==source:
     OUTPUTS.append(str(source));break
    try:
     os.link(source,target)
     try:source.unlink()
     except Exception:target.unlink();raise
     OUTPUTS.append(str(target));break
    except FileExistsError:continue
   else:raise UserError('Troppi nomi duplicati.')
 elif action=='files.rename':
  prefix=nonempty(p,'prefix')
  if '/' in prefix or '\\' in prefix or prefix in ['.','..']:raise UserError('Prefisso non valido.')
  for i,path in enumerate(paths,int(number(p,'start',1,0))):
   fake=Path(path).with_name(prefix+f'-{i:04}'+Path(path).suffix);publish(path,fake,Path(path).suffix[1:],suffix='')
 else:raise UserError('Questa azione viene gestita dall’interfaccia nativa.')
 return OUTPUTS

if __name__=='__main__':
 reply_fd=globals().get('RESPONSE_FD')
 if reply_fd is None:reply_fd=os.dup(1);os.dup2(2,1)
 code=0
 try:
  req=json.loads(Path(sys.argv[1]).read_text());execute(req);response={'ok':True,'outputs':OUTPUTS}
 except Exception as e:
  response={'ok':False,'error':str(e),'outputs':OUTPUTS};code=1
 sys.stdout.flush()
 with os.fdopen(reply_fd,'w',encoding='utf-8') as reply:json.dump(response,reply,ensure_ascii=False);reply.write('\n')
 sys.exit(code)
