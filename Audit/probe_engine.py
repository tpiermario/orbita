import sys,os,json,tempfile,io,zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'Engine'));import engine as E
from PIL import Image
import pymupdf as F
E.RES=ROOT/'Orbita.app/Contents/Resources'
D=Path(tempfile.mkdtemp(prefix='orbita-audit-')); results=[]
def probe(name,fn):
 try:
  detail=fn(); results.append({'name':name,'evidence':detail});print(name,repr(detail),flush=True)
 except Exception as e:print('PROBE ERROR',name,type(e).__name__,str(e)[:200],flush=True)
def op(a,p,params={}):
 E.OUTPUTS.clear();E.execute({'action':a,'paths':[str(x) for x in p],'params':params});return [Path(x) for x in E.OUTPUTS]
def pic(name='a.png',size=(40,30),color='red'):
 p=D/name;Image.new('RGBA',size,color).save(p);return p
png=pic()
def gifloop():
 p=D/'once.gif';Image.new('RGB',(10,10),'red').save(p,save_all=True,append_images=[Image.new('RGB',(10,10),'blue')],duration=100)
 q=op('image.to.gif',[p])[0];return {'input_loop':Image.open(p).info.get('loop'),'output_loop':Image.open(q).info.get('loop')}
probe('GIF once becomes infinite',gifloop)
def redact():
 p=pic('box.png',(20,20),'white');q=op('image.redact',[p],{'x':2,'y':2,'width':1,'height':1})[0];im=Image.open(q).convert('RGB');return {'requested_pixels':1,'black_pixels':sum(x==(0,0,0) for x in im.getdata())}
probe('Image redaction rectangle off by one',redact)
def ico():
 p=pic('tiny.png',(8,8));q=op('image.to.ico',[p])[0]
 try: im=Image.open(q);return {'bytes':q.stat().st_size,'sizes':list(im.ico.sizes())}
 except Exception as e:return {'bytes':q.stat().st_size,'decode_error':str(e)}
probe('Tiny ICO invalid',ico)
def compress():
 p=pic('panorama.png',(10000,10));
 try:op('image.compress',[p],{'mb':.001});return 'success'
 except Exception as e:return {'error':str(e),'feasible_800x1_bytes':len(E.image_bytes(Image.new('RGB',(800,1)),'jpg',quality=10,optimize=True))}
probe('Panorama compression stops prematurely',compress)
def icc():
 im=Image.new('RGB',(20,20));im.info['icc_profile']=Path('/System/Library/ColorSync/Profiles/Display P3.icc').read_bytes();p=D/'icc.png';im.save(p,icc_profile=im.info['icc_profile']);q=op('image.to.jpg',[p])[0];return {'input_has_profile':bool(Image.open(p).info.get('icc_profile')),'output_has_profile':bool(Image.open(q).info.get('icc_profile'))}
probe('JPEG drops ICC profile',icc)
def alpha_mask():
 p=pic('transparent.png',(8,8),(255,0,0,0));old=E.mask_subject;E.mask_subject=lambda im:Image.new('L',im.size,255)
 try:q=op('image.removebg',[p])[0];return {'input_alpha':0,'output_alpha':Image.open(q).getpixel((0,0))[3]}
 finally:E.mask_subject=old
probe('Remove background restores already transparent pixels',alpha_mask)
def meta():
 p=D/'orient.jpg';ex=Image.Exif();ex[274]=6;Image.new('RGB',(40,20),'red').save(p,exif=ex);q=op('image.metadata',[p])[0];return json.loads(q.read_text())
probe('Read metadata reports transformed dimensions and removes orientation',meta)
def makepdf(name='doc.pdf'):
 d=F.open();d.new_page(width=300,height=400).insert_text((50,50),'Hello World');p=D/name;d.save(p);return p
pdf=makepdf()
def unicodepdf():
 q=op('pdf.addtext',[pdf],{'text':'Ελληνικά 中文 €','x':30,'y':120})[0];return F.open(q)[0].get_text()
probe('PDF text loses Unicode',unicodepdf)
def extractalpha():
 d=F.open();page=d.new_page();im=pic('alpha.png',(10,10),(255,0,0,80));page.insert_image(F.Rect(0,0,20,20),filename=im);p=D/'alpha.pdf';d.save(p);q=op('pdf.images',[p])[0];return {'mode':Image.open(q).mode,'has_transparency':Image.open(q).info.get('transparency') is not None}
probe('PDF image extraction loses soft mask',extractalpha)
def pptaspect():
 from pptx import Presentation
 d=F.open();d.new_page(width=300,height=400);d.new_page(width=400,height=300);p=D/'mixed.pdf';d.save(p);q=op('pdf.to.pptx',[p])[0];r=Presentation(q);s=r.slides[1].shapes[0];return {'original_aspect':400/300,'output_aspect':s.width/s.height}
probe('PDF to PPTX stretches mixed orientation pages',pptaspect)
def ocrduplicate():
 old=E.vision;E.vision=lambda *a,**k:[{'text':'Hello World','rect':[.1,.1,.4,.05]}]
 try:q=op('pdf.ocr',[pdf])[0];return F.open(q)[0].get_text()
 finally:E.vision=old
probe('Searchable PDF OCR duplicates existing text',ocrduplicate)
def rotation_text():
 d=F.open(pdf);d[0].set_rotation(90);p=D/'rotated.pdf';d.save(p);q=op('pdf.numbers',[p])[0];d=F.open(q);return {'page_visual_rect':list(d[0].rect),'inserted_1': '1' in d[0].get_text(),'text':d[0].get_text()}
probe('Number rotated PDF uses wrong coordinates',rotation_text)
def pdfbounds():
 q=op('pdf.addtext',[pdf],{'text':'INVISIBLE','x':10000,'y':10000})[0];return {'output_exists':q.exists(),'text':F.open(q)[0].get_text()}
probe('PDF accepts offpage insertion with silent loss',pdfbounds)
def compresspdf():
 q=op('pdf.compress',[pdf])[0];return {'before':pdf.stat().st_size,'after':q.stat().st_size}
probe('PDF compress enlarges simple document',compresspdf)
def outline():
 d=F.open(pdf);d.set_toc([[1,'Chapter',1]]);p=D/'bookmarked.pdf';d.save(p);q=op('pdf.merge',[p,p])[0];return {'input_toc':F.open(p).get_toc(),'output_toc':F.open(q).get_toc()}
probe('PDF merge discards navigation bookmarks',outline)
def active_sheet():
 from openpyxl import Workbook
 p=D/'sheets.xlsx';w=Workbook();w.active.append(['FIRST']);w.create_sheet('Second').append(['SECOND']);w.active=1;w.save(p);q=op('sheet.to.csv',[p])[0];return q.read_text(encoding='utf-8-sig')
probe('Spreadsheet exports active sheet instead of documented first',active_sheet)
def formulas():
 from openpyxl import Workbook
 p=D/'formula.xlsx';w=Workbook();w.active.append([1,2,'=SUM(A1:B1)']);w.save(p);q=op('sheet.to.json',[p])[0];return json.loads(q.read_text())
probe('Spreadsheet formula without cache becomes null',formulas)
def apostrophe():
 p=D/'values.csv';p.write_text('-10,+391234,=literal,@name\n');q=op('sheet.to.xlsx',[p])[0];r=op('sheet.to.csv',[q])[0];return r.read_text(encoding='utf-8-sig')
probe('Literal spreadsheet text gains apostrophe',apostrophe)
def odsvalue():
 from odf.opendocument import OpenDocumentSpreadsheet
 from odf.table import Table,TableRow,TableCell
 d=OpenDocumentSpreadsheet();t=Table(name='Data');r=TableRow();r.addElement(TableCell(valuetype='float',value=42));t.addElement(r);d.spreadsheet.addElement(t);p=D/'numeric.ods';d.save(str(p));return E.read_sheet(p)
probe('ODS numeric value without display text becomes empty',odsvalue)
def odsrepeat():
 from odf.opendocument import OpenDocumentSpreadsheet
 from odf.table import Table,TableRow,TableCell
 from odf.text import P
 d=OpenDocumentSpreadsheet();t=Table(name='Data');r=TableRow(numberrowsrepeated=10001);c=TableCell(valuetype='string');c.addElement(P(text='x'));r.addElement(c);t.addElement(r);d.spreadsheet.addElement(t);p=D/'repeated.ods';d.save(str(p));return {'expected':10001,'actual':len(E.read_sheet(p))}
probe('ODS repetition truncates data silently',odsrepeat)
def odscovered():
 from odf.opendocument import OpenDocumentSpreadsheet
 from odf.table import Table,TableRow,TableCell,CoveredTableCell
 from odf.text import P
 d=OpenDocumentSpreadsheet();t=Table(name='Data');r=TableRow();c=TableCell(valuetype='string',numbercolumnsspanned=2);c.addElement(P(text='A'));r.addElement(c);r.addElement(CoveredTableCell());c=TableCell(valuetype='string');c.addElement(P(text='C'));r.addElement(c);t.addElement(r);d.spreadsheet.addElement(t);p=D/'merged.ods';d.save(str(p));return E.read_sheet(p)
probe('ODS covered cells shift following columns left',odscovered)
def longcell():
 from openpyxl import load_workbook
 p=D/'long.json';p.write_text(json.dumps([['x'*40000]]));q=op('sheet.to.xlsx',[p])[0];return {'input_len':40000,'output_len':len(load_workbook(q).active['A1'].value)}
probe('XLSX conversion silently truncates long cells',longcell)
def txt16():
 p=D/'utf16.txt';p.write_text('Ciao mondo',encoding='utf-16')
 try:op('document.to.docx',[p]);return 'success'
 except Exception as e:return str(e)
probe('TextEdit UTF16 text rejected',txt16)
def resource():
 p=D/'local.md';p.write_text('![red](a.png)');q=op('document.to.pdf',[p])[0];d=F.open(q);return {'embedded_images':sum(len(page.get_images()) for page in d),'text':d[0].get_text()}
probe('Markdown relative images omitted',resource)
def standalonehtml():
 from docx import Document
 d=Document();d.add_picture(str(png));p=D/'image.docx';d.save(p);q=op('document.to.html',[p])[0];text=q.read_text();import re;src=re.findall(r'<img[^>]*src="([^"]+)',text);return {'src':src,'exists':[ (q.parent/x).exists() for x in src if not x.startswith('data:')]}
probe('Office HTML conversion publishes without images',standalonehtml)
def renamedot():
 p=D/'README';p.write_text('test');q=op('files.rename',[p],{'prefix':'Copy'})[0];return q.name
probe('Extensionless renamed copies gain trailing dot',renamedot)
def renameidempotent():
 p=D/'File-0001.txt';p.write_text('test');q=op('files.renameOriginals',[p],{'prefix':'File','start':1})[0];return q.name
probe('Rename already matching name adds collision suffix',renameidempotent)
def stripcustom():
 from docx import Document
 from docx.oxml import OxmlElement
 p=D/'custom.docx';Document().save(p)
 with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
 parts['docProps/custom.xml']=b'<?xml version="1.0"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Owner"><vt:lpwstr>PRIVATE_AUTHOR</vt:lpwstr></property></Properties>'
 parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'</Types>',b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
 parts['_rels/.rels']=parts['_rels/.rels'].replace(b'</Relationships>',b'<Relationship Id="rIdAuditCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
 with zipfile.ZipFile(p,'w') as z:
  for n,data in parts.items():z.writestr(n,data)
 q=op('files.strip',[p])[0]
 with zipfile.ZipFile(q) as z:return {'private_metadata_remains':b'PRIVATE_AUTHOR' in z.read('docProps/custom.xml')}
probe('Office metadata removal leaves custom properties',stripcustom)
(D/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));(ROOT/'Audit/engine-probes.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));print('FIXTURES',D)
